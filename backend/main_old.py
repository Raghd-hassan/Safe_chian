import os
import io
from fastapi import FastAPI, Depends, HTTPException, status, Query, Request
from fastapi.security import OAuth2PasswordBearer, HTTPAuthorizationCredentials, HTTPBearer
from sqlalchemy.orm import Session
from sqlalchemy import func
from database import SessionLocal
import models
from passlib.context import CryptContext
from schemas import UserCreate, UserLogin, ResetPassword, VerifyCode, UserUpdate
from datetime import datetime, timedelta, time
from jose import JWTError, jwt
import random
from pydantic import BaseModel
from typing import List, Optional
from fastapi.middleware.cors import CORSMiddleware
from collections import defaultdict
from threading import Lock
import openpyxl
from fastapi.responses import StreamingResponse
from fpdf import FPDF
import arabic_reshaper
from bidi.algorithm import get_display

# ============================================================
# ✅ الإعدادات الأساسية
# ============================================================
SECRET_KEY = os.getenv("SECRET_KEY", "s3cr3t_k3y_s4f3ch41n_2024_change_me")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 7  # 7 أيام

app = FastAPI(title="SafeChain API", version="1.0.0")

# ============================================================
# ✅ CORS
# ============================================================
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ============================================================
# ✅ الأمان
# ============================================================
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="login", auto_error=False)
security = HTTPBearer(auto_error=False)

# ============================================================
# ✅ Rate Limiting
# ============================================================
rate_limit_storage = defaultdict(list)
rate_limit_lock = Lock()

def check_rate_limit(ip: str, max_attempts: int = 5, window_seconds: int = 60):
    now = datetime.now()
    with rate_limit_lock:
        rate_limit_storage[ip] = [
            t for t in rate_limit_storage[ip] if (now - t).total_seconds() < window_seconds
        ]
        if len(rate_limit_storage[ip]) >= max_attempts:
            raise HTTPException(status_code=429, detail="محاولات كثيرة جداً. حاول بعد قليل")
        rate_limit_storage[ip].append(now)

# ============================================================
# ✅ تخزين OTP المؤقت
# ============================================================
otp_storage: dict[str, dict] = {}

def store_otp(email: str, otp: str):
    otp_storage[email] = {"code": otp, "expires_at": datetime.now() + timedelta(minutes=10)}

def verify_otp(email: str, code: str) -> bool:
    entry = otp_storage.get(email)
    if not entry: return False
    if datetime.now() > entry["expires_at"]:
        otp_storage.pop(email, None)
        return False
    return entry["code"] == code

def clear_otp(email: str):
    otp_storage.pop(email, None)

# ============================================================
# ✅ اتصال قاعدة البيانات
# ============================================================
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# ============================================================
# ✅ دوال JWT والأمان
# ============================================================
def create_access_token(data: dict, expires_delta: timedelta | None = None):
    to_encode = data.copy()
    expire = datetime.utcnow() + (expires_delta or timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
):
    if not credentials:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="لم يتم توفير رمز المصادقة", headers={"WWW-Authenticate": "Bearer"})

    token = credentials.credentials
    credentials_exception = HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="رمز المصادقة غير صالح أو منتهي", headers={"WWW-Authenticate": "Bearer"})

    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = payload.get("sub")
        if user_id is None: raise credentials_exception
        user_id = int(user_id)
    except JWTError as e:
        print(f"JWT Error: {e}")
        raise credentials_exception

    Users = models.Base.classes.users
    user = db.query(Users).filter(Users.ID == user_id).first()
    if user is None: raise credentials_exception

    jwt_role = payload.get("role")
    if jwt_role and jwt_role != user.role:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="تم تحديث صلاحياتك، سجّل دخولك مجدداً")

    return user

def require_admin(current_user=Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="غير مصرح لك بهذا الإجراء")
    return current_user

# ============================================================
# ✅ فلترة مشتركة
# ============================================================
def apply_filters(query, model, start_date, end_date, search):
    if start_date:
        try:
            start_dt = datetime.combine(datetime.strptime(start_date, "%Y-%m-%d").date(), time.min)
            query = query.filter(model.full_date >= start_dt)
        except Exception: pass
    if end_date:
        try:
            end_dt = datetime.combine(datetime.strptime(end_date, "%Y-%m-%d").date(), time.max)
            query = query.filter(model.full_date <= end_dt)
        except Exception: pass
    if search:
        if hasattr(model, 'shipment_number'):
            query = query.filter(model.shipment_number.ilike(f"%{search}%"))
    return query

# ============================================================
# ✅ نماذج Pydantic
# ============================================================
class EmailRequest(BaseModel):
    email: str

class StartTripRequest(BaseModel):
    truck_id: int
    product_id: int | None = None
    group_id: int | None = None

class CreateAlertRequest(BaseModel):
    shipment_number: str
    reason: str
    latitude: float = 0.0
    longitude: float = 0.0

class CreateReportRequest(BaseModel):
    shipment_number: str
    temperature: float
    humidity: int
    gas_1: float
    gas_2: float
    door_condition: str
    latitude: float | None = None
    longitude: float | None = None

class SecurityLimitsRequest(BaseModel):
    responsible_name: str
    truck_id: int
    products: List[str]
    min_temp: float
    max_temp: float
    min_humidity: int
    max_humidity: int
    door_open_duration: int
    door_open_count: int

class UpdateRoleRequest(BaseModel):
    role: str

def reshape_arabic(text: str) -> str:
    if not text: return ""
    reshaped = arabic_reshaper.reshape(text)
    return get_display(reshaped)

# ============================================================
# ✅ التحقق من التوكن والمصادقة
# ============================================================
@app.get("/verify-token")
async def verify_token(current_user=Depends(get_current_user)):
    return {"valid": True, "user_id": current_user.ID, "role": current_user.role}

@app.post("/register")
def register_user(request: Request, user: UserCreate, db: Session = Depends(get_db)):
    check_rate_limit(request.client.host, max_attempts=3, window_seconds=60)
    Users = models.Base.classes.users
    exist = db.query(Users).filter(Users.email == user.email).first()
    if exist: raise HTTPException(status_code=400, detail="الإيميل مسجل مسبقاً")

    hashed_password = pwd_context.hash(user.password)
    new_user = Users(full_name=user.full_name, email=user.email, phone=user.phone, password=hashed_password, role="user")
    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    otp = str(random.randint(1000, 9999))
    store_otp(user.email, otp)
    print(f"🔐 OTP for {user.email}: {otp}")
    return {"message": "تم إنشاء الحساب", "otp": otp}

@app.post("/send-reset-code")
def send_reset_code(request: Request, data: EmailRequest, db: Session = Depends(get_db)):
    check_rate_limit(request.client.host, max_attempts=3, window_seconds=60)
    Users = models.Base.classes.users
    user = db.query(Users).filter(Users.email == data.email).first()
    if not user: raise HTTPException(status_code=404, detail="الإيميل غير موجود")

    otp = str(random.randint(1000, 9999))
    store_otp(data.email, otp)
    print(f"🔐 RESET OTP for {data.email}: {otp}")
    return {"message": "تم إرسال رمز التحقق"}

@app.post("/verify-email")
def verify_email(data: VerifyCode):
    if not verify_otp(data.email, data.code): raise HTTPException(status_code=400, detail="رمز غير صحيح أو منتهي الصلاحية")
    clear_otp(data.email)
    return {"message": "تم التفعيل"}

@app.post("/login")
def login_user(request: Request, credentials: UserLogin, db: Session = Depends(get_db)):
    check_rate_limit(request.client.host, max_attempts=5, window_seconds=60)
    Users = models.Base.classes.users
    user = db.query(Users).filter(Users.email == credentials.email).first()
    if not user or not pwd_context.verify(credentials.password, user.password):
        raise HTTPException(status_code=401, detail="بيانات غير صحيحة")

    access_token = create_access_token(data={"sub": str(user.ID), "role": user.role})
    return {"access_token": access_token, "token_type": "bearer", "role": user.role, "user_id": user.ID, "full_name": user.full_name, "message": "تم تسجيل الدخول"}

@app.post("/reset-password")
def reset_password(data: ResetPassword, db: Session = Depends(get_db)):
    Users = models.Base.classes.users
    user = db.query(Users).filter(Users.email == data.email).first()
    if not user: raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if not verify_otp(data.email, data.code): raise HTTPException(status_code=400, detail="رمز غير صحيح أو منتهي الصلاحية")
    user.password = pwd_context.hash(data.new_password)
    db.commit()
    clear_otp(data.email)
    return {"message": "تم تغيير كلمة المرور"}

# ============================================================
# ✅ الملف الشخصي
# ============================================================
@app.get("/profile")
def get_user_profile(current_user=Depends(get_current_user)):
    return {"full_name": current_user.full_name, "email": current_user.email, "phone": current_user.phone, "role": current_user.role}

@app.put("/profile")
def update_profile(data: UserUpdate, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Users = models.Base.classes.users
    user_in_db = db.query(Users).filter(Users.ID == current_user.ID).first()
    if not user_in_db: raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if data.full_name: user_in_db.full_name = data.full_name
    if data.email:
        email_exists = db.query(Users).filter(Users.email == data.email, Users.ID != current_user.ID).first()
        if email_exists: raise HTTPException(status_code=400, detail="الإيميل مستخدم من قبل شخص آخر")
        user_in_db.email = data.email
    if data.phone: user_in_db.phone = data.phone
    if data.password: user_in_db.password = pwd_context.hash(data.password)
    db.commit()
    return {"message": "تم تحديث البيانات بنجاح"}

# ============================================================
# ✅ الشاحنات والرحلات
# ============================================================
@app.get("/trucks")
def get_trucks(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks
    return db.query(Trucks).filter(Trucks.user_id == current_user.ID).all()

@app.post("/start-trip")
def start_trip(trip_data: StartTripRequest, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    TripCounter = models.Base.classes.trip_counter
    Trips = models.Base.classes.trips
    Trucks = models.Base.classes.trucks

    truck = db.query(Trucks).filter(Trucks.ID == trip_data.truck_id, Trucks.user_id == current_user.ID).first()
    if not truck: raise HTTPException(status_code=404, detail="الشاحنة غير موجودة أو لا تخصك")

    active_trip = db.query(Trips).filter(Trips.truck_id == trip_data.truck_id, Trips.status == "active").first()
    if active_trip: raise HTTPException(status_code=400, detail="توجد رحلة نشطة بالفعل لهذه الشاحنة")

    counter = db.query(TripCounter).filter(TripCounter.truck_id == trip_data.truck_id, TripCounter.user_id == current_user.ID).first()
    if counter: counter.last_sequence += 1
    else:
        counter = TripCounter(truck_id=trip_data.truck_id, user_id=current_user.ID, last_sequence=1)
        db.add(counter)
    db.flush()

    seq = str(counter.last_sequence).zfill(4)
    shipment_number = f"SHP-{trip_data.truck_id}-{seq}"
    new_trip = Trips(truck_id=trip_data.truck_id, shipment_number=shipment_number, status="active", start_time=datetime.now(), user_id=current_user.ID, product_id=trip_data.product_id, group_id=trip_data.group_id)

    try:
        db.add(new_trip); db.commit(); db.refresh(new_trip)
        return {"status": "success", "shipment_number": shipment_number}
    except Exception as e:
        db.rollback(); raise HTTPException(status_code=500, detail=str(e))

@app.post("/end-trip")
def end_trip(shipment_number: str = Query(..., description="رقم الشحنة"), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trips = models.Base.classes.trips
    Reports = models.Base.classes.reports
    trip = db.query(Trips).filter(Trips.shipment_number == shipment_number, Trips.user_id == current_user.ID).first()
    if not trip: raise HTTPException(status_code=404, detail="الرحلة غير موجودة أو لا تخصك")
    if trip.status == "completed": raise HTTPException(status_code=400, detail="الرحلة منتهية مسبقاً")

    last_report = db.query(Reports).filter(Reports.truck_id == trip.truck_id).order_by(Reports.full_date.desc()).first()
    if last_report:
        trip.end_temperature = getattr(last_report, 'temperature', None)
        trip.end_humidity = getattr(last_report, 'humidity', None)

    trip.status = "completed"; trip.end_time = datetime.now()
    try:
        db.commit()
        return {"status": "success", "message": "تم إنهاء الرحلة بنجاح", "shipment_number": shipment_number, "end_time": trip.end_time.strftime("%Y-%m-%d %H:%M:%S")}
    except Exception as e:
        db.rollback(); raise HTTPException(status_code=500, detail=f"حدث خطأ: {str(e)}")

@app.get("/truck-live-data")
def get_truck_live_data(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks
    Reports = models.Base.classes.reports
    truck = db.query(Trucks).filter(Trucks.user_id == current_user.ID).first()
    if not truck: return {}
    report = db.query(Reports).filter(Reports.truck_id == truck.ID).order_by(Reports.full_date.desc()).first()
    if not report: return {}
    return {"shipment_number": report.shipment_number, "temperature": report.temperature, "humidity": report.humidity, "gas_1": getattr(report, "mq9_gas", None), "gas_2": getattr(report, "mq135_gas", None), "door_condition": getattr(report, "door_condition", None), "latitude": report.latitude, "longitude": report.longitude, "last_updated": str(report.full_date)}

# ============================================================
# ✅ التقارير والتنبيهات
# ============================================================
@app.get("/reports")
def get_reports(truck_id: int = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Reports = models.Base.classes.reports; Trucks = models.Base.classes.trucks
    query = db.query(Reports).join(Trucks, Reports.truck_id == Trucks.ID).filter(Trucks.user_id == current_user.ID)
    if truck_id: query = query.filter(Reports.truck_id == truck_id)
    query = apply_filters(query, Reports, start_date, end_date, search)
    total = query.count(); skip = (page - 1) * limit; results = query.order_by(Reports.full_date.desc()).offset(skip).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "data": [{"id": r.ID, "latitude": r.latitude, "longitude": r.longitude, "shipment_number": r.shipment_number, "temperature": r.temperature, "humidity": r.humidity, "door_condition": getattr(r, 'door_condition', None), "gas_1": getattr(r, 'mq9_gas', None), "gas_2": getattr(r, 'mq135_gas', None), "full_date": r.full_date.strftime("%Y-%m-%d %H:%M:%S")} for r in results]}

@app.post("/save-report")
def save_report_periodic(data: CreateReportRequest, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trips = models.Base.classes.trips; Reports = models.Base.classes.reports
    trip = db.query(Trips).filter(Trips.shipment_number == data.shipment_number, Trips.user_id == current_user.ID, Trips.status == "active").first()
    if not trip: raise HTTPException(status_code=404, detail="الرحلة غير موجودة أو منتهية")
    if data.door_condition not in ("OPEN", "CLOSED"): raise HTTPException(status_code=400, detail="door_condition يجب أن يكون OPEN أو CLOSED")
    new_report = Reports(truck_id=trip.truck_id, shipment_number=data.shipment_number, temperature=data.temperature, humidity=data.humidity, mq9_gas=data.gas_1, mq135_gas=data.gas_2, door_condition=data.door_condition, latitude=data.latitude or 0.0, longitude=data.longitude or 0.0, full_date=datetime.now())
    db.add(new_report); db.commit()
    return {"status": "success", "message": "تم حفظ التقرير الدوري"}

@app.get("/alerts")
def get_alerts(truck_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips
    query = db.query(AlertsModel).filter(AlertsModel.user_id == current_user.ID)
    if truck_id: query = query.outerjoin(Trips, AlertsModel.shipment_number == Trips.shipment_number).filter(Trips.truck_id == truck_id)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); results = query.order_by(AlertsModel.full_date.desc()).all()
    return {"data": [{"id": r.ID, "shipment_number": r.shipment_number, "type": r.alert_reason, "duration": r.duration, "date": r.full_date.strftime("%Y-%m-%d %H:%M:%S"), "latitude": r.latitude, "longitude": r.longitude} for r in results]}

@app.post("/save-alert")
def save_alert(data: CreateAlertRequest, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips
    trip = db.query(Trips).filter(Trips.shipment_number == data.shipment_number, Trips.user_id == current_user.ID).first()
    if not trip: raise HTTPException(status_code=404, detail="رقم الشحنة غير موجود أو لا يخصك")
    existing_alert = db.query(AlertsModel).filter(AlertsModel.shipment_number == data.shipment_number).order_by(AlertsModel.full_date.desc()).first()
    if existing_alert:
        if (datetime.now() - existing_alert.full_date).total_seconds() < 60: return {"status": "ignored", "message": "تم تجاهل التنبيه المكرر"}
    new_alert = AlertsModel(shipment_number=data.shipment_number, alert_reason=data.reason, duration="جاري...", user_id=current_user.ID, full_date=datetime.now(), latitude=data.latitude, longitude=data.longitude)
    db.add(new_alert); db.commit()
    return {"status": "success", "message": "تم حفظ التنبيه بنجاح"}

# ============================================================
# ✅ المنتجات وحدود الأمان
# ============================================================
@app.get("/products")
def get_products(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Products = models.Base.classes.products; Groups = models.Base.classes.groups
    results = db.query(Products, Groups).outerjoin(Groups, Groups.ID == Products.group_id).all()
    return [{"id": p.ID, "product_name": p.product_name, "min_temperature": p.min_temperature, "max_temperature": p.max_temperature, "min_humidity": p.min_humidity, "max_humidity": p.max_humidity, "group_id": p.group_id or None, "group_min_temperature": g.min_temperature if g else None, "group_max_temperature": g.max_temperature if g else None, "group_min_humidity": g.min_humidity if g else None, "group_max_humidity": g.max_humidity if g else None} for p, g in results]

@app.post("/save-security-limits")
def save_security_limits(data: SecurityLimitsRequest, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks
    if data.min_temp >= data.max_temp: raise HTTPException(status_code=400, detail="الحرارة الدنيا يجب أن تكون أقل من العليا")
    if data.min_humidity >= data.max_humidity: raise HTTPException(status_code=400, detail="الرطوبة الدنيا يجب أن تكون أقل من العليا")
    truck = db.query(Trucks).filter(Trucks.ID == data.truck_id, Trucks.user_id == current_user.ID).first()
    if not truck: raise HTTPException(status_code=404, detail="الشاحنة غير موجودة")
    for attr, val in [('min_temp', data.min_temp), ('max_temp', data.max_temp), ('min_humidity', data.min_humidity), ('max_humidity', data.max_humidity), ('door_open_duration', data.door_open_duration), ('door_open_count', data.door_open_count)]:
        if hasattr(truck, attr): setattr(truck, attr, val)
    db.commit()
    return {"status": "success", "message": "تم حفظ حدود الأمان بنجاح"}

# ============================================================
# ✅ لوحة تحكم المستخدم
# ============================================================
@app.get("/dashboard")
def get_dashboard(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks; Reports = models.Base.classes.reports
    user_trucks = db.query(Trucks).filter(Trucks.user_id == current_user.ID).all()
    truck_ids = [t.ID for t in user_trucks]
    if not truck_ids: return {"total_trucks": 0, "compliant_trucks": 0, "non_compliant_trucks": 0}

    subquery = db.query(Reports.truck_id, func.max(Reports.full_date).label('max_date')).filter(Reports.truck_id.in_(truck_ids)).group_by(Reports.truck_id).subquery()
    latest_reports = db.query(Reports).join(subquery, (Reports.truck_id == subquery.c.truck_id) & (Reports.full_date == subquery.c.max_date)).all()
    report_map = {r.truck_id: r for r in latest_reports}; truck_map = {t.ID: t for t in user_trucks}

    compliant_trucks = 0; non_compliant_trucks = 0
    for truck_id in truck_ids:
        truck = truck_map[truck_id]; report = report_map.get(truck_id); is_compliant = True
        if report:
            min_temp = getattr(truck, 'min_temp', None); max_temp = getattr(truck, 'max_temp', None)
            if min_temp is not None and max_temp is not None:
                try:
                    if not (float(min_temp) <= float(report.temperature) <= float(max_temp)): is_compliant = False
                except: pass
            if is_compliant:
                min_hum = getattr(truck, 'min_humidity', None); max_hum = getattr(truck, 'max_humidity', None)
                if min_hum is not None and max_hum is not None:
                    try:
                        if not (float(min_hum) <= float(report.humidity) <= float(max_hum)): is_compliant = False
                    except: pass
        if is_compliant: compliant_trucks += 1
        else: non_compliant_trucks += 1
    return {"total_trucks": len(user_trucks), "compliant_trucks": compliant_trucks, "non_compliant_trucks": non_compliant_trucks}

@app.get("/sensors")
def get_sensors(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks; Reports = models.Base.classes.reports
    user_trucks = db.query(Trucks).filter(Trucks.user_id == current_user.ID).all()
    truck_ids = [t.ID for t in user_trucks]; sensors_list = []
    if not truck_ids: return sensors_list

    subquery = db.query(Reports.truck_id, func.max(Reports.full_date).label('max_date')).filter(Reports.truck_id.in_(truck_ids)).group_by(Reports.truck_id).subquery()
    latest_reports = db.query(Reports).join(subquery, (Reports.truck_id == subquery.c.truck_id) & (Reports.full_date == subquery.c.max_date)).all()
    report_map = {r.truck_id: r for r in latest_reports}

    for truck in user_trucks:
        report = report_map.get(truck.ID)
        temp_status = hum_status = gas1_status = gas2_status = door_status = "غير معروف"
        if report:
            temp_status = hum_status = gas1_status = gas2_status = door_status = "شغال"
            min_temp = getattr(truck, 'min_temp', None); max_temp = getattr(truck, 'max_temp', None)
            if min_temp is not None and max_temp is not None:
                try:
                    if not (float(min_temp) <= float(report.temperature) <= float(max_temp)): temp_status = "تنبيه"
                except: temp_status = "خطأ"
            min_hum = getattr(truck, 'min_humidity', None); max_hum = getattr(truck, 'max_humidity', None)
            if min_hum is not None and max_hum is not None:
                try:
                    if not (float(min_hum) <= float(report.humidity) <= float(max_hum)): hum_status = "تنبيه"
                except: hum_status = "خطأ"
            if getattr(report, 'mq9_gas', None) is None: gas1_status = "غير معروف"
            if getattr(report, 'mq135_gas', None) is None: gas2_status = "غير معروف"
            if getattr(report, 'door_condition', None) == "OPEN": door_status = "تنبيه"
        sensors_list.extend([
            {"name": "حساس الحرارة", "serial_number": f"SN-TEMP-{truck.ID:03d}", "status": temp_status},
            {"name": "حساس الرطوبة", "serial_number": f"SN-HUM-{truck.ID:03d}", "status": hum_status},
            {"name": "حساس الغاز MQ9", "serial_number": f"SN-GAS1-{truck.ID:03d}", "status": gas1_status},
            {"name": "حساس الغاز MQ135", "serial_number": f"SN-GAS2-{truck.ID:03d}", "status": gas2_status},
            {"name": "حساس الباب", "serial_number": f"SN-DOOR-{truck.ID:03d}", "status": door_status},
        ])
    return sensors_list

@app.get("/updates")
def get_updates(start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips; Trucks = models.Base.classes.trucks; Products = models.Base.classes.products
    query = db.query(AlertsModel).filter(AlertsModel.user_id == current_user.ID)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); results = query.order_by(AlertsModel.full_date.desc()).all()
    updates_list = []
    for alert in results:
        truck_id_val = None; product_name = ""; temp_range = ""
        if alert.shipment_number:
            trip = db.query(Trips).filter(Trips.shipment_number == alert.shipment_number).first()
            if trip:
                truck_id_val = trip.truck_id
                if trip.product_id:
                    product = db.query(Products).filter(Products.ID == trip.product_id).first()
                    if product: product_name = product.product_name
                if truck_id_val:
                    truck = db.query(Trucks).filter(Trucks.ID == truck_id_val).first()
                    if truck:
                        min_t = getattr(truck, 'min_temp', None); max_t = getattr(truck, 'max_temp', None)
                        if min_t is not None and max_t is not None: temp_range = f"{min_t}°C إلى {max_t}°C"
        updates_list.append({"truck_id": truck_id_val, "shipment_number": alert.shipment_number or "-", "alert_reason": alert.alert_reason or "-", "product": product_name, "admin": current_user.full_name, "temp_range": temp_range, "date": alert.full_date.strftime("%d/%m/%Y %I:%M %p") if alert.full_date else ""})
    return updates_list

# ============================================================
# ✅ تصدير ملفات المستخدم العادي
# ============================================================
@app.get("/export-excel-alerts")
def export_excel_alerts(truck_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips
    query = db.query(AlertsModel).filter(AlertsModel.user_id == current_user.ID)
    if truck_id: query = query.outerjoin(Trips, AlertsModel.shipment_number == Trips.shipment_number).filter(Trips.truck_id == truck_id)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); results = query.order_by(AlertsModel.full_date.desc()).all()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "التنبيهات"
    ws.append(["التاريخ", "سبب التنبيه", "رقم الشحنة", "مدة التنبيه", "خط العرض", "خط الطول"])
    for r in results: ws.append([r.full_date.strftime("%Y-%m-%d %H:%M:%S") if r.full_date else "", r.alert_reason or "", r.shipment_number or "", r.duration or "", str(r.latitude) if r.latitude else "", str(r.longitude) if r.longitude else ""])
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="alerts_export.xlsx"'})

@app.get("/export-pdf-alerts")
def export_pdf_alerts(truck_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips
    query = db.query(AlertsModel).filter(AlertsModel.user_id == current_user.ID)
    if truck_id: query = query.outerjoin(Trips, AlertsModel.shipment_number == Trips.shipment_number).filter(Trips.truck_id == truck_id)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); results = query.order_by(AlertsModel.full_date.desc()).all()
    pdf = FPDF(orientation='L', unit='mm', format='A4'); pdf.add_page(); pdf.set_font("Helvetica", size=12)
    pdf.set_font_size(18); pdf.cell(0, 10, reshape_arabic("تقرير التنبيهات - SafeChain"), ln=True, align='C'); pdf.ln(10)
    pdf.set_font_size(10); pdf.set_fill_color(27, 67, 50); pdf.set_text_color(255, 255, 255)
    col_widths = [50, 90, 40, 30, 40, 40]; headers_list = ["التاريخ", "سبب التنبيه", "رقم الشحنة", "المدة", "خط العرض", "خط الطول"]
    for i, header in enumerate(headers_list): pdf.cell(col_widths[i], 10, reshape_arabic(header), border=1, align='C', fill=True)
    pdf.ln(); pdf.set_text_color(0, 0, 0); pdf.set_font_size(9)
    for r in results:
        pdf.cell(col_widths[0], 8, r.full_date.strftime("%Y-%m-%d %H:%M") if r.full_date else "-", border=1, align='C')
        pdf.cell(col_widths[1], 8, reshape_arabic(r.alert_reason or "-"), border=1, align='C')
        pdf.cell(col_widths[2], 8, r.shipment_number or "-", border=1, align='C')
        pdf.cell(col_widths[3], 8, reshape_arabic(r.duration or "-"), border=1, align='C')
        pdf.cell(col_widths[4], 8, str(r.latitude) if r.latitude else "-", border=1, align='C')
        pdf.cell(col_widths[5], 8, str(r.longitude) if r.longitude else "-", border=1, align='C')
        pdf.ln()
    stream = io.BytesIO(); pdf.output(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename="alerts_export.pdf"'})

@app.get("/export-excel-reports")
def export_excel_reports(truck_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Reports = models.Base.classes.reports; Trucks = models.Base.classes.trucks
    query = db.query(Reports).join(Trucks, Reports.truck_id == Trucks.ID).filter(Trucks.user_id == current_user.ID)
    if truck_id: query = query.filter(Reports.truck_id == truck_id)
    query = apply_filters(query, Reports, start_date, end_date, search); results = query.order_by(Reports.full_date.desc()).all()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "التقارير"
    ws.append(["التاريخ", "الحرارة", "الرطوبة", "حساس الغاز 1 (MQ9)", "حساس الغاز 2 (MQ135)", "حالة الباب", "رقم الشحنة", "خط العرض", "خط الطول"])
    for r in results: ws.append([r.full_date.strftime("%Y-%m-%d %H:%M:%S") if r.full_date else "", str(r.temperature) if r.temperature else "", str(r.humidity) if r.humidity else "", str(getattr(r, 'mq9_gas', '')) if getattr(r, 'mq9_gas', None) is not None else "", str(getattr(r, 'mq135_gas', '')) if getattr(r, 'mq135_gas', None) is not None else "", getattr(r, 'door_condition', '') or "", r.shipment_number or "", str(r.latitude) if r.latitude else "", str(r.longitude) if r.longitude else ""])
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="reports_export.xlsx"'})

@app.get("/export-pdf-reports")
def export_pdf_reports(truck_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Reports = models.Base.classes.reports; Trucks = models.Base.classes.trucks
    query = db.query(Reports).join(Trucks, Reports.truck_id == Trucks.ID).filter(Trucks.user_id == current_user.ID)
    if truck_id: query = query.filter(Reports.truck_id == truck_id)
    query = apply_filters(query, Reports, start_date, end_date, search); results = query.order_by(Reports.full_date.desc()).all()
    pdf = FPDF(orientation='L', unit='mm', format='A4'); pdf.add_page(); pdf.set_font("Helvetica", size=12)
    pdf.set_font_size(18); pdf.cell(0, 10, reshape_arabic("تقرير القراءات - SafeChain"), ln=True, align='C'); pdf.ln(10)
    pdf.set_font_size(9); pdf.set_fill_color(27, 67, 50); pdf.set_text_color(255, 255, 255)
    col_widths = [40, 25, 25, 30, 30, 25, 35, 40, 40]; headers_list = ["التاريخ", "الحرارة", "الرطوبة", "غاز 1 (MQ9)", "غاز 2 (MQ135)", "الباب", "رقم الشحنة", "خط العرض", "خط الطول"]
    for i, header in enumerate(headers_list): pdf.cell(col_widths[i], 10, reshape_arabic(header), border=1, align='C', fill=True)
    pdf.ln(); pdf.set_text_color(0, 0, 0); pdf.set_font_size(8)
    for r in results:
        pdf.cell(col_widths[0], 8, r.full_date.strftime("%Y-%m-%d %H:%M") if r.full_date else "-", border=1, align='C')
        pdf.cell(col_widths[1], 8, str(r.temperature) if r.temperature else "-", border=1, align='C')
        pdf.cell(col_widths[2], 8, str(r.humidity) if r.humidity else "-", border=1, align='C')
        pdf.cell(col_widths[3], 8, str(getattr(r, 'mq9_gas', '-')) if getattr(r, 'mq9_gas', None) is not None else "-", border=1, align='C')
        pdf.cell(col_widths[4], 8, str(getattr(r, 'mq135_gas', '-')) if getattr(r, 'mq135_gas', None) is not None else "-", border=1, align='C')
        pdf.cell(col_widths[5], 8, reshape_arabic(getattr(r, 'door_condition', '-') or "-"), border=1, align='C')
        pdf.cell(col_widths[6], 8, r.shipment_number or "-", border=1, align='C')
        pdf.cell(col_widths[7], 8, str(r.latitude) if r.latitude else "-", border=1, align='C')
        pdf.cell(col_widths[8], 8, str(r.longitude) if r.longitude else "-", border=1, align='C')
        pdf.ln()
    stream = io.BytesIO(); pdf.output(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename="reports_export.pdf"'})

@app.get("/export-excel-sensors")
def export_excel_sensors(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks; Reports = models.Base.classes.reports
    user_trucks = db.query(Trucks).filter(Trucks.user_id == current_user.ID).all(); truck_ids = [t.ID for t in user_trucks]
    if not truck_ids: raise HTTPException(status_code=404, detail="لا توجد شاحنات مسجلة")
    subquery = db.query(Reports.truck_id, func.max(Reports.full_date).label('max_date')).filter(Reports.truck_id.in_(truck_ids)).group_by(Reports.truck_id).subquery()
    latest_reports = db.query(Reports).join(subquery, (Reports.truck_id == subquery.c.truck_id) & (Reports.full_date == subquery.c.max_date)).all()
    report_map = {r.truck_id: r for r in latest_reports}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "الحساسات"; ws.append(["اسم الحساس", "الرقم التسلسلي", "الحالة"])
    for truck in user_trucks:
        report = report_map.get(truck.ID)
        temp_status = "غير معروف"
        if report:
            temp_status = "شغال"; min_temp = getattr(truck, 'min_temp', None); max_temp = getattr(truck, 'max_temp', None)
            if min_temp is not None and max_temp is not None:
                try:
                    if not (float(min_temp) <= float(report.temperature) <= float(max_temp)): temp_status = "تنبيه"
                except: pass
        ws.append(["حساس الحرارة", f"SN-TEMP-{truck.ID:03d}", temp_status])
        hum_status = "غير معروف"
        if report: 
            hum_status = "شغال"; min_hum = getattr(truck, 'min_humidity', None); max_hum = getattr(truck, 'max_humidity', None)
            if min_hum is not None and max_hum is not None:
                try:
                    if not (float(min_hum) <= float(report.humidity) <= float(max_hum)): hum_status = "تنبيه"
                except: pass
        ws.append(["حساس الرطوبة", f"SN-HUM-{truck.ID:03d}", hum_status])
        gas1_status = "غير معروف"
        if report and getattr(report, 'mq9_gas', None) is not None: gas1_status = "شغال"
        ws.append(["حساس الغاز MQ9", f"SN-GAS1-{truck.ID:03d}", gas1_status])
        gas2_status = "غير معروف"
        if report and getattr(report, 'mq135_gas', None) is not None: gas2_status = "شغال"
        ws.append(["حساس الغاز MQ135", f"SN-GAS2-{truck.ID:03d}", gas2_status])
        door_status = "غير معروف"
        if report:
            door_status = "شغال"
            if getattr(report, 'door_condition', None) == "OPEN": door_status = "تنبيه"
        ws.append(["حساس الباب", f"SN-DOOR-{truck.ID:03d}", door_status])
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="sensors_export.xlsx"'})

@app.get("/export-pdf-sensors")
def export_pdf_sensors(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    Trucks = models.Base.classes.trucks; Reports = models.Base.classes.reports
    user_trucks = db.query(Trucks).filter(Trucks.user_id == current_user.ID).all(); truck_ids = [t.ID for t in user_trucks]
    if not truck_ids: raise HTTPException(status_code=404, detail="لا توجد شاحنات مسجلة")
    subquery = db.query(Reports.truck_id, func.max(Reports.full_date).label('max_date')).filter(Reports.truck_id.in_(truck_ids)).group_by(Reports.truck_id).subquery()
    latest_reports = db.query(Reports).join(subquery, (Reports.truck_id == subquery.c.truck_id) & (Reports.full_date == subquery.c.max_date)).all()
    report_map = {r.truck_id: r for r in latest_reports}
    pdf = FPDF(orientation='L', unit='mm', format='A4'); pdf.add_page(); pdf.set_font("Helvetica", size=12)
    pdf.set_font_size(18); pdf.cell(0, 10, reshape_arabic("تقرير فحص الحساسات - SafeChain"), ln=True, align='C'); pdf.ln(10)
    pdf.set_font_size(10); pdf.set_fill_color(27, 67, 50); pdf.set_text_color(255, 255, 255)
    col_widths = [80, 60, 40]; headers_list = ["اسم الحساس", "الرقم التسلسلي", "الحالة"]
    for i, header in enumerate(headers_list): pdf.cell(col_widths[i], 10, reshape_arabic(header), border=1, align='C', fill=True)
    pdf.ln(); pdf.set_text_color(0, 0, 0); pdf.set_font_size(9)
    for truck in user_trucks:
        report = report_map.get(truck.ID); rows = []
        temp_status = "غير معروف"
        if report:
            temp_status = "شغال"; min_temp = getattr(truck, 'min_temp', None); max_temp = getattr(truck, 'max_temp', None)
            if min_temp is not None and max_temp is not None:
                try:
                    if not (float(min_temp) <= float(report.temperature) <= float(max_temp)): temp_status = "تنبيه"
                except: pass
        rows.append(["حساس الحرارة", f"SN-TEMP-{truck.ID:03d}", temp_status])
        hum_status = "غير معروف"
        if report:
            hum_status = "شغال"; min_hum = getattr(truck, 'min_humidity', None); max_hum = getattr(truck, 'max_humidity', None)
            if min_hum is not None and max_hum is not None:
                try:
                    if not (float(min_hum) <= float(report.humidity) <= float(max_hum)): hum_status = "تنبيه"
                except: pass
        rows.append(["حساس الرطوبة", f"SN-HUM-{truck.ID:03d}", hum_status])
        gas1_status = "غير معروف"
        if report and getattr(report, 'mq9_gas', None) is not None: gas1_status = "شغال"
        rows.append(["حساس الغاز MQ9", f"SN-GAS1-{truck.ID:03d}", gas1_status])
        gas2_status = "غير معروف"
        if report and getattr(report, 'mq135_gas', None) is not None: gas2_status = "شغال"
        rows.append(["حساس الغاز MQ135", f"SN-GAS2-{truck.ID:03d}", gas2_status])
        door_status = "غير معروف"
        if report:
            door_status = "شغال"
            if getattr(report, 'door_condition', None) == "OPEN": door_status = "تنبيه"
        rows.append(["حساس الباب", f"SN-DOOR-{truck.ID:03d}", door_status])
        for row in rows:
            pdf.cell(col_widths[0], 8, reshape_arabic(row[0]), border=1, align='C')
            pdf.cell(col_widths[1], 8, row[1], border=1, align='C')
            pdf.cell(col_widths[2], 8, reshape_arabic(row[2]), border=1, align='C')
            pdf.ln()
    stream = io.BytesIO(); pdf.output(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename="sensors_export.pdf"'})

@app.get("/export-excel-updates")
def export_excel_updates(start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips; Trucks = models.Base.classes.trucks; Products = models.Base.classes.products
    query = db.query(AlertsModel).filter(AlertsModel.user_id == current_user.ID)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); results = query.order_by(AlertsModel.full_date.desc()).all()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "سجل التحديثات"
    ws.append(["رقم الشاحنة", "رقم الشحنة", "سبب التنبيه", "المنتج", "المسؤول", "تعديل الحدود", "التاريخ"])
    for alert in results:
        truck_id_val = None; product_name = ""; temp_range = ""
        if alert.shipment_number:
            trip = db.query(Trips).filter(Trips.shipment_number == alert.shipment_number).first()
            if trip:
                truck_id_val = trip.truck_id
                if trip.product_id:
                    product = db.query(Products).filter(Products.ID == trip.product_id).first()
                    if product: product_name = product.product_name
                if truck_id_val:
                    truck = db.query(Trucks).filter(Trucks.ID == truck_id_val).first()
                    if truck:
                        min_t = getattr(truck, 'min_temp', None); max_t = getattr(truck, 'max_temp', None)
                        if min_t is not None and max_t is not None: temp_range = f"{min_t}°C إلى {max_t}°C"
        ws.append([truck_id_val or "-", alert.shipment_number or "-", alert.alert_reason or "-", product_name or "-", current_user.full_name, temp_range or "-", alert.full_date.strftime("%d/%m/%Y %I:%M %p") if alert.full_date else ""])
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="updates_export.xlsx"'})

@app.get("/export-pdf-updates")
def export_pdf_updates(start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips; Trucks = models.Base.classes.trucks; Products = models.Base.classes.products
    query = db.query(AlertsModel).filter(AlertsModel.user_id == current_user.ID)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); results = query.order_by(AlertsModel.full_date.desc()).all()
    pdf = FPDF(orientation='L', unit='mm', format='A4'); pdf.add_page(); pdf.set_font("Helvetica", size=12)
    pdf.set_font_size(18); pdf.cell(0, 10, reshape_arabic("سجل التحديثات - SafeChain"), ln=True, align='C'); pdf.ln(10)
    pdf.set_font_size(9); pdf.set_fill_color(27, 67, 50); pdf.set_text_color(255, 255, 255)
    col_widths = [25, 40, 65, 35, 40, 40, 45]; headers_list = ["الشاحنة", "الشحنة", "سبب التنبيه", "المنتج", "المسؤول", "تعديل الحدود", "التاريخ"]
    for i, header in enumerate(headers_list): pdf.cell(col_widths[i], 10, reshape_arabic(header), border=1, align='C', fill=True)
    pdf.ln(); pdf.set_text_color(0, 0, 0); pdf.set_font_size(8)
    for alert in results:
        truck_id_val = None; product_name = ""; temp_range = ""
        if alert.shipment_number:
            trip = db.query(Trips).filter(Trips.shipment_number == alert.shipment_number).first()
            if trip:
                truck_id_val = trip.truck_id
                if trip.product_id:
                    product = db.query(Products).filter(Products.ID == trip.product_id).first()
                    if product: product_name = product.product_name
                if truck_id_val:
                    truck = db.query(Trucks).filter(Trucks.ID == truck_id_val).first()
                    if truck:
                        min_t = getattr(truck, 'min_temp', None); max_t = getattr(truck, 'max_temp', None)
                        if min_t is not None and max_t is not None: temp_range = f"{min_t} C to {max_t} C"
        pdf.cell(col_widths[0], 8, str(truck_id_val or "-"), border=1, align='C')
        pdf.cell(col_widths[1], 8, alert.shipment_number or "-", border=1, align='C')
        pdf.cell(col_widths[2], 8, reshape_arabic(alert.alert_reason or "-"), border=1, align='C')
        pdf.cell(col_widths[3], 8, reshape_arabic(product_name or "-"), border=1, align='C')
        pdf.cell(col_widths[4], 8, reshape_arabic(current_user.full_name), border=1, align='C')
        pdf.cell(col_widths[5], 8, temp_range or "-", border=1, align='C')
        pdf.cell(col_widths[6], 8, alert.full_date.strftime("%Y-%m-%d %H:%M") if alert.full_date else "-", border=1, align='C')
        pdf.ln()
    stream = io.BytesIO(); pdf.output(stream); stream.seek(0)
    return StreamingResponse(stream, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename="updates_export.pdf"'})

# ============================================================
# ✅ Admin - إدارة المستخدمين
# ============================================================
@app.get("/admin/users")
def get_all_users(search: str = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Users = models.Base.classes.users; query = db.query(Users)
    if search: query = query.filter(Users.full_name.ilike(f"%{search}%") | Users.email.ilike(f"%{search}%"))
    total = query.count(); skip = (page - 1) * limit; users = query.offset(skip).limit(limit).all()
    return {"total": total, "page": page, "limit": limit, "data": [{"id": u.ID, "name": u.full_name, "email": u.email, "phone": u.phone, "role": u.role} for u in users]}

@app.delete("/admin/users/{user_id}")
def delete_user(user_id: int, db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Users = models.Base.classes.users; user = db.query(Users).filter(Users.ID == user_id).first()
    if not user: raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if user.ID == current_user.ID: raise HTTPException(status_code=400, detail="لا يمكنك حذف نفسك")
    db.delete(user); db.commit()
    return {"message": "تم حذف المستخدم"}

@app.put("/admin/users/{user_id}/role")
def update_user_role(user_id: int, data: UpdateRoleRequest, db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Users = models.Base.classes.users; user = db.query(Users).filter(Users.ID == user_id).first()
    if not user: raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if data.role not in ["admin", "user"]: raise HTTPException(status_code=400, detail="قيمة role غير صحيحة")
    user.role = data.role; db.commit()
    return {"message": "تم تحديث الصلاحية"}

# ============================================================
# ✅ Admin - لوحة تحكم الأدمن الشاملة (تم الإصلاح)
# ============================================================
@app.get("/admin/dashboard")
def admin_dashboard(db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Users = models.Base.classes.users; Trucks = models.Base.classes.trucks; Trips = models.Base.classes.trips; Reports = models.Base.classes.reports; AlertsModel = models.Base.classes.alerts
    total_users = db.query(Users).count(); total_trucks = db.query(Trucks).count(); active_trips = db.query(Trips).filter(Trips.status == "active").count(); total_reports = db.query(Reports).count(); total_alerts = db.query(AlertsModel).count()

    compliant_trucks = 0; non_compliant_trucks = 0
    if total_trucks > 0:
        all_trucks = db.query(Trucks).all(); truck_ids = [t.ID for t in all_trucks]
        if truck_ids:
            subquery = db.query(Reports.truck_id, func.max(Reports.full_date).label('max_date')).filter(Reports.truck_id.in_(truck_ids)).group_by(Reports.truck_id).subquery()
            latest_reports = db.query(Reports).join(subquery, (Reports.truck_id == subquery.c.truck_id) & (Reports.full_date == subquery.c.max_date)).all()
            report_map = {r.truck_id: r for r in latest_reports}; truck_map = {t.ID: t for t in all_trucks}

            for truck_id in truck_ids:
                truck = truck_map[truck_id]; report = report_map.get(truck_id); is_compliant = True
                if report:
                    has_temp_limits = getattr(truck, 'min_temp', None) is not None and getattr(truck, 'max_temp', None) is not None
                    has_hum_limits = getattr(truck, 'min_humidity', None) is not None and getattr(truck, 'max_humidity', None) is not None
                    if not has_temp_limits and not has_hum_limits:
                        is_compliant = False
                    else:
                        if has_temp_limits:
                            try:
                                if not (float(truck.min_temp) <= float(report.temperature) <= float(truck.max_temp)): is_compliant = False
                            except: is_compliant = False
                        if is_compliant and has_hum_limits:
                            try:
                                if not (float(truck.min_humidity) <= float(report.humidity) <= float(truck.max_humidity)): is_compliant = False
                            except: is_compliant = False
                else:
                    is_compliant = False
                if is_compliant: compliant_trucks += 1
                else: non_compliant_trucks += 1

    return {"total_users": total_users, "total_trucks": total_trucks, "active_trips": active_trips, "total_reports": total_reports, "total_alerts": total_alerts, "compliant_trucks": compliant_trucks, "non_compliant_trucks": non_compliant_trucks}

# ============================================================
# ✅ Admin - التقارير والتنبيهات والرحلات
# ============================================================
@app.get("/admin/reports")
def admin_get_reports(user_id: int = Query(None), truck_id: int = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Reports = models.Base.classes.reports; Trucks = models.Base.classes.trucks
    query = db.query(Reports).join(Trucks, Reports.truck_id == Trucks.ID)
    if user_id: query = query.filter(Trucks.user_id == user_id)
    if truck_id: query = query.filter(Reports.truck_id == truck_id)
    query = apply_filters(query, Reports, start_date, end_date, search); total = query.count(); skip = (page - 1) * limit; results = query.order_by(Reports.full_date.desc()).offset(skip).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "data": [{"id": r.ID, "truck_id": r.truck_id, "shipment_number": r.shipment_number, "temperature": r.temperature, "humidity": r.humidity, "gas_1": getattr(r, "mq9_gas", None), "gas_2": getattr(r, "mq135_gas", None), "door_condition": getattr(r, "door_condition", None), "latitude": r.latitude, "longitude": r.longitude, "full_date": r.full_date.strftime("%Y-%m-%d %H:%M:%S")} for r in results]}

@app.get("/admin/alerts")
def admin_get_alerts(user_id: int = Query(None), truck_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips; Trucks = models.Base.classes.trucks
    query = db.query(AlertsModel)
    if user_id: query = query.filter(AlertsModel.user_id == user_id)
    if truck_id: query = query.outerjoin(Trips, AlertsModel.shipment_number == Trips.shipment_number).outerjoin(Trucks, Trips.truck_id == Trucks.ID).filter(Trucks.ID == truck_id)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); total = query.count(); skip = (page - 1) * limit; results = query.order_by(AlertsModel.full_date.desc()).offset(skip).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "data": [{"id": r.ID, "user_id": r.user_id, "shipment_number": r.shipment_number, "type": r.alert_reason, "duration": r.duration, "date": r.full_date.strftime("%Y-%m-%d %H:%M:%S"), "latitude": r.latitude, "longitude": r.longitude} for r in results]}

@app.get("/admin/trips")
def admin_get_trips(user_id: int = Query(None), truck_id: int = Query(None), status: str = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Trips = models.Base.classes.trips; query = db.query(Trips)
    if user_id: query = query.filter(Trips.user_id == user_id)
    if truck_id: query = query.filter(Trips.truck_id == truck_id)
    if status: query = query.filter(Trips.status == status)
    if search: query = query.filter(Trips.shipment_number.ilike(f"%{search}%"))
    if start_date:
        try: query = query.filter(Trips.start_time >= datetime.combine(datetime.strptime(start_date, "%Y-%m-%d").date(), time.min))
        except: pass
    if end_date:
        try: query = query.filter(Trips.start_time <= datetime.combine(datetime.strptime(end_date, "%Y-%m-%d").date(), time.max))
        except: pass
    total = query.count(); skip = (page - 1) * limit; results = query.order_by(Trips.start_time.desc()).offset(skip).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "data": [{"id": t.ID, "user_id": t.user_id, "truck_id": t.truck_id, "shipment_number": t.shipment_number, "status": t.status, "start_time": t.start_time.strftime("%Y-%m-%d %H:%M:%S") if t.start_time else None, "end_time": t.end_time.strftime("%Y-%m-%d %H:%M:%S") if getattr(t, "end_time", None) else None, "product_id": getattr(t, "product_id", None), "group_id": getattr(t, "group_id", None)} for t in results]}

@app.get("/admin/trucks")
def admin_get_trucks(user_id: int = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Trucks = models.Base.classes.trucks; query = db.query(Trucks)
    if user_id: query = query.filter(Trucks.user_id == user_id)
    total = query.count(); skip = (page - 1) * limit; results = query.offset(skip).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "data": [{"id": t.ID, "user_id": t.user_id, "min_temp": getattr(t, "min_temp", None), "max_temp": getattr(t, "max_temp", None), "min_humidity": getattr(t, "min_humidity", None), "max_humidity": getattr(t, "max_humidity", None), "door_open_duration": getattr(t, "door_open_duration", None), "door_open_count": getattr(t, "door_open_count", None)} for t in results]}

@app.get("/admin/sensors")
def admin_get_sensors(user_id: int = Query(None), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    Trucks = models.Base.classes.trucks; Reports = models.Base.classes.reports
    query = db.query(Trucks)
    if user_id: query = query.filter(Trucks.user_id == user_id)
    user_trucks = query.all(); truck_ids = [t.ID for t in user_trucks]
    if not truck_ids: return []
    subquery = db.query(Reports.truck_id, func.max(Reports.full_date).label("max_date")).filter(Reports.truck_id.in_(truck_ids)).group_by(Reports.truck_id).subquery()
    latest_reports = db.query(Reports).join(subquery, (Reports.truck_id == subquery.c.truck_id) & (Reports.full_date == subquery.c.max_date)).all()
    report_map = {r.truck_id: r for r in latest_reports}; sensors_list = []
    for truck in user_trucks:
        report = report_map.get(truck.ID); temp_status = hum_status = gas1_status = gas2_status = door_status = "غير معروف"
        if report:
            temp_status = hum_status = gas1_status = gas2_status = door_status = "شغال"
            min_temp = getattr(truck, "min_temp", None); max_temp = getattr(truck, "max_temp", None)
            if min_temp is not None and max_temp is not None:
                try:
                    if not (float(min_temp) <= float(report.temperature) <= float(max_temp)): temp_status = "تنبيه"
                except: temp_status = "خطأ"
            min_hum = getattr(truck, "min_humidity", None); max_hum = getattr(truck, "max_humidity", None)
            if min_hum is not None and max_hum is not None:
                try:
                    if not (float(min_hum) <= float(report.humidity) <= float(max_hum)): hum_status = "تنبيه"
                except: hum_status = "خطأ"
            if getattr(report, "mq9_gas", None) is None: gas1_status = "غير معروف"
            if getattr(report, "mq135_gas", None) is None: gas2_status = "غير معروف"
            if getattr(report, "door_condition", None) == "OPEN": door_status = "تنبيه"
        sensors_list.extend([
            {"truck_id": truck.ID, "user_id": truck.user_id, "name": "حساس الحرارة", "serial_number": f"SN-TEMP-{truck.ID:03d}", "status": temp_status},
            {"truck_id": truck.ID, "user_id": truck.user_id, "name": "حساس الرطوبة", "serial_number": f"SN-HUM-{truck.ID:03d}", "status": hum_status},
            {"truck_id": truck.ID, "user_id": truck.user_id, "name": "حساس الغاز MQ9", "serial_number": f"SN-GAS1-{truck.ID:03d}", "status": gas1_status},
            {"truck_id": truck.ID, "user_id": truck.user_id, "name": "حساس الغاز MQ135", "serial_number": f"SN-GAS2-{truck.ID:03d}", "status": gas2_status},
            {"truck_id": truck.ID, "user_id": truck.user_id, "name": "حساس الباب", "serial_number": f"SN-DOOR-{truck.ID:03d}", "status": door_status},
        ])
    return sensors_list

@app.get("/admin/updates")
def admin_get_updates(user_id: int = Query(None), start_date: str = Query(None), end_date: str = Query(None), search: str = Query(None), page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), current_user=Depends(require_admin)):
    AlertsModel = models.Base.classes.alerts; Trips = models.Base.classes.trips; Trucks = models.Base.classes.trucks; Products = models.Base.classes.products; Users = models.Base.classes.users
    query = db.query(AlertsModel)
    if user_id: query = query.filter(AlertsModel.user_id == user_id)
    query = apply_filters(query, AlertsModel, start_date, end_date, search); total = query.count(); skip = (page - 1) * limit; results = query.order_by(AlertsModel.full_date.desc()).offset(skip).limit(limit).all()
    updates_list = []
    for alert in results:
        truck_id_val = None; product_name = ""; temp_range = ""; owner_name = ""
        owner = db.query(Users).filter(Users.ID == alert.user_id).first()
        if owner: owner_name = owner.full_name
        if alert.shipment_number:
            trip = db.query(Trips).filter(Trips.shipment_number == alert.shipment_number).first()
            if trip:
                truck_id_val = trip.truck_id
                if trip.product_id:
                    product = db.query(Products).filter(Products.ID == trip.product_id).first()
                    if product: product_name = product.product_name
                if truck_id_val:
                    truck = db.query(Trucks).filter(Trucks.ID == truck_id_val).first()
                    if truck:
                        min_t = getattr(truck, "min_temp", None); max_t = getattr(truck, "max_temp", None)
                        if min_t is not None and max_t is not None: temp_range = f"{min_t}°C إلى {max_t}°C"
        updates_list.append({"truck_id": truck_id_val, "user_id": alert.user_id, "owner_name": owner_name, "shipment_number": alert.shipment_number or "-", "alert_reason": alert.alert_reason or "-", "product": product_name, "temp_range": temp_range, "date": alert.full_date.strftime("%d/%m/%Y %I:%M %p") if alert.full_date else ""})
    return {"page": page, "limit": limit, "total": total, "data": updates_list}

